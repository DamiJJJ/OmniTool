import io
import os
import logging
from xml.sax.saxutils import escape as xml_escape

from flask import (
    Blueprint,
    render_template,
    request,
    flash,
    redirect,
    url_for,
    send_file,
)
from flask_wtf import FlaskForm
from wtforms import FileField, SelectField, StringField, SubmitField
from wtforms.validators import DataRequired, Length
from PIL import Image

from extensions import limiter

Image.MAX_IMAGE_PIXELS = 50_000_000

logger = logging.getLogger(__name__)

_MAGIC_CHECKS = [
    (lambda h: h[:3] == b'\xff\xd8\xff'),
    (lambda h: h[:8] == b'\x89PNG\r\n\x1a\n'),
    (lambda h: h[:6] in (b'GIF87a', b'GIF89a')),
    (lambda h: h[:2] == b'BM'),
    (lambda h: h[:4] in (b'II*\x00', b'MM\x00*')),
    (lambda h: h[:4] == b'RIFF' and h[8:12] == b'WEBP'),
]


def _is_valid_image(stream: io.BytesIO) -> bool:
    header = stream.read(12)
    stream.seek(0)
    return any(check(header) for check in _MAGIC_CHECKS)


def _is_pdf_bytes(data: bytes) -> bool:
    return data[:5] == b"%PDF-"


def _is_zip_bytes(data: bytes) -> bool:
    return data[:4] in (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")


conversion_bp = Blueprint("conversion", __name__, template_folder="../templates")

# Supported formats for the dropdowns (Pillow compatibility)
# Note: "JPG" is mapped to "JPEG" for PIL compatibility
IMAGE_FORMATS = [
    ("AUTO", "Auto Detect"),
    ("JPEG", "JPEG/JPG"),
    ("PNG", "PNG"),
    ("WEBP", "WebP"),
    ("GIF", "GIF"),
    ("BMP", "BMP"),
    ("TIFF", "TIFF/TIF"),
]

# Max file size in bytes (10 MB)
MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_MERGE_FILES = 10
MAX_PDF_PAGES = 500
MAX_TEXT_CHARS = 500_000


# --- Forms ---
class ConversionForm(FlaskForm):
    source_format = SelectField(
        "Convert From", choices=IMAGE_FORMATS, validators=[DataRequired()]
    )
    target_format = SelectField(
        "Convert To",
        choices=IMAGE_FORMATS[1:],
        validators=[DataRequired()],
    )
    file = FileField("Image File", validators=[DataRequired()])
    submit = SubmitField("Convert Image")


class SingleFileForm(FlaskForm):
    """Generic single-file upload form (CSRF + submit)."""
    file = FileField("File", validators=[DataRequired()])
    submit = SubmitField("Convert")


class MergePDFForm(FlaskForm):
    """Multi-file form: validation handled manually via request.files.getlist."""
    submit = SubmitField("Merge PDFs")


class SplitPDFForm(FlaskForm):
    file = FileField("PDF File", validators=[DataRequired()])
    page_ranges = StringField(
        "Page ranges",
        validators=[DataRequired(), Length(max=200)],
    )
    submit = SubmitField("Split PDF")


# --- Helpers ---
def _read_upload(file_storage):
    """Read uploaded file fully into bytes, enforcing MAX_FILE_SIZE.

    Returns bytes or raises ValueError.
    """
    file_storage.seek(0, os.SEEK_END)
    size = file_storage.tell()
    file_storage.seek(0)
    if size == 0:
        raise ValueError("Uploaded file is empty.")
    if size > MAX_FILE_SIZE:
        raise ValueError(
            f"File size exceeds the limit of {MAX_FILE_SIZE // (1024 * 1024)} MB."
        )
    return file_storage.read()


def _parse_page_ranges(spec: str, max_page: int) -> list[int]:
    """Parse e.g. '1-3,5,7-9' against a 1-indexed page count.

    Returns a sorted list of unique page numbers (1-indexed). Raises ValueError on bad input.
    """
    if not spec or len(spec) > 200:
        raise ValueError("Page ranges are required (max 200 characters).")
    pages: set[int] = set()
    for raw in spec.split(","):
        part = raw.strip()
        if not part:
            continue
        if "-" in part:
            a_str, b_str = part.split("-", 1)
            a_str, b_str = a_str.strip(), b_str.strip()
            if not (a_str.isdigit() and b_str.isdigit()):
                raise ValueError(f"Invalid range: '{part}'.")
            a, b = int(a_str), int(b_str)
            if a < 1 or b < 1 or a > b or b > max_page:
                raise ValueError(
                    f"Range '{part}' is out of bounds (document has {max_page} pages)."
                )
            pages.update(range(a, b + 1))
        else:
            if not part.isdigit():
                raise ValueError(f"Invalid page number: '{part}'.")
            n = int(part)
            if n < 1 or n > max_page:
                raise ValueError(
                    f"Page {n} is out of bounds (document has {max_page} pages)."
                )
            pages.add(n)
    if not pages:
        raise ValueError("No valid pages selected.")
    return sorted(pages)


_UNICODE_FONT_NAME: str | None = None


def _get_pdf_font() -> str:
    """Try to register a system Unicode TTF for proper non-ASCII rendering.

    Falls back to built-in Helvetica (Latin-1 only) if no system font is found.
    The result is cached after first call.
    """
    global _UNICODE_FONT_NAME
    if _UNICODE_FONT_NAME is not None:
        return _UNICODE_FONT_NAME
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/Library/Fonts/Microsoft/Verdana.ttf",
        "C:/Windows/Fonts/arial.ttf",
    ]
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        for path in candidates:
            if os.path.exists(path):
                pdfmetrics.registerFont(TTFont("AppUnicode", path))
                _UNICODE_FONT_NAME = "AppUnicode"
                return _UNICODE_FONT_NAME
    except Exception:
        logger.exception("Failed to register Unicode TTF; falling back to Helvetica.")
    _UNICODE_FONT_NAME = "Helvetica"
    return _UNICODE_FONT_NAME


def _text_to_pdf_stream(text: str) -> io.BytesIO:
    """Render plain text into a paginated PDF using reportlab Platypus."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    if len(text) > MAX_TEXT_CHARS:
        raise ValueError(
            f"Text exceeds the maximum length of {MAX_TEXT_CHARS:,} characters."
        )

    font_name = _get_pdf_font()
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        name="OmniBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=11,
        leading=14,
    )

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title="OmniTool Conversion",
    )
    story = []
    # Split on blank lines into paragraphs; preserve line breaks within.
    paragraphs = text.replace("\r\n", "\n").replace("\r", "\n").split("\n\n")
    for paragraph in paragraphs:
        safe = xml_escape(paragraph).replace("\n", "<br/>")
        if not safe.strip():
            story.append(Spacer(1, 8))
            continue
        story.append(Paragraph(safe, body_style))
        story.append(Spacer(1, 6))
    if not story:
        story.append(Paragraph("&nbsp;", body_style))
    doc.build(story)
    out.seek(0)
    return out


def _docx_to_text(data: bytes) -> str:
    """Extract plain text from a DOCX byte string."""
    from docx import Document

    document = Document(io.BytesIO(data))
    lines: list[str] = []
    for para in document.paragraphs:
        lines.append(para.text)
    # Tables → tab-separated rows.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.replace("\t", " ") for cell in row.cells]
            lines.append("\t".join(cells))
    return "\n".join(lines).strip()


def _safe_send(stream: io.BytesIO, mimetype: str, filename: str):
    return send_file(
        stream,
        mimetype=mimetype,
        as_attachment=True,
        download_name=filename,
        max_age=0,
    )


# --- Conversion Center Index ---
@conversion_bp.route("/convert", methods=["GET"])
def conversion_index():
    """Renders the main conversion hub with module selection."""
    return render_template("conversion_index.html")


# --- Image Conversion ---
@conversion_bp.route("/convert/image", methods=["GET", "POST"])
@limiter.limit("30 per minute;200 per hour", methods=["POST"])
def image_conversion():
    """Renders the image conversion form (GET) and handles conversion logic (POST)."""
    form = ConversionForm()

    if request.method == "POST":
        if form.validate_on_submit():
            file = request.files.get("file")
            target_format = form.target_format.data

            if not file:
                flash("No file part or file not selected.", "danger")
                return redirect(url_for("conversion.image_conversion"))

            file.seek(0, os.SEEK_END)
            file_size = file.tell()
            file.seek(0)

            if file_size > MAX_FILE_SIZE:
                flash(
                    f"File size exceeds the limit of {MAX_FILE_SIZE / (1024 * 1024):.0f}MB.",
                    "danger",
                )
                return redirect(url_for("conversion.image_conversion"))

            try:
                image_stream = io.BytesIO(file.read())
                if not _is_valid_image(image_stream):
                    flash("Uploaded file is not a recognised image format.", "danger")
                    return redirect(url_for("conversion.image_conversion"))
                img = Image.open(image_stream)
                img.verify()
                image_stream.seek(0)
                img = Image.open(image_stream)

                pil_format = target_format

                if target_format in ["JPEG", "JPG"]:
                    pil_format = "JPEG"
                    if img.mode == "RGBA":
                        img.load()
                        background = Image.new("RGB", img.size, (255, 255, 255))
                        background.paste(img, mask=img.split()[3])
                        img = background

                elif target_format == "GIF":
                    img = img.convert("P") if img.mode != "P" else img

                output_stream = io.BytesIO()

                img.save(output_stream, format=pil_format)
                output_stream.seek(0)

                filename = (
                    f"converted_image.{target_format.lower().replace('jpeg', 'jpg')}"
                )

                mime_type = f"image/{target_format.lower()}"
                if pil_format == "JPEG":
                    mime_type = "image/jpeg"
                elif pil_format == "TIFF":
                    mime_type = "image/tiff"
                elif pil_format == "BMP":
                    mime_type = "image/bmp"

                return send_file(
                    output_stream,
                    mimetype=mime_type,
                    as_attachment=True,
                    download_name=filename,
                )

            except Exception as e:
                logger.exception("Image conversion failed")
                flash(f"An error occurred during conversion: {e}", "danger")
                return redirect(url_for("conversion.image_conversion"))

        flash(
            "Form validation failed. Please check your selection and try again.",
            "danger",
        )
        return redirect(url_for("conversion.image_conversion"))

    return render_template("image_conversion.html", form=form)


# --- Document & PDF Tools Hub ---
@conversion_bp.route("/convert/documents", methods=["GET"])
def document_conversion():
    """Renders the Document & PDF Tools hub."""
    return render_template("document_conversion.html")


# --- Merge PDFs ---
@conversion_bp.route("/convert/documents/merge-pdf", methods=["GET", "POST"])
@limiter.limit("15 per minute;100 per hour", methods=["POST"])
def merge_pdf():
    form = MergePDFForm()

    if request.method == "POST":
        if not form.validate_on_submit():
            flash("Form validation failed. Please try again.", "danger")
            return redirect(url_for("conversion.merge_pdf"))

        files = request.files.getlist("files")
        files = [f for f in files if f and f.filename]
        if len(files) < 2:
            flash("Please select at least 2 PDF files to merge.", "danger")
            return redirect(url_for("conversion.merge_pdf"))
        if len(files) > MAX_MERGE_FILES:
            flash(f"You can merge at most {MAX_MERGE_FILES} files at once.", "danger")
            return redirect(url_for("conversion.merge_pdf"))

        try:
            from pypdf import PdfWriter, PdfReader

            writer = PdfWriter()
            total_pages = 0
            for f in files:
                data = _read_upload(f)
                if not _is_pdf_bytes(data):
                    raise ValueError(f"'{f.filename}' is not a valid PDF file.")
                reader = PdfReader(io.BytesIO(data), strict=False)
                if reader.is_encrypted:
                    raise ValueError(f"'{f.filename}' is encrypted; please decrypt it first.")
                total_pages += len(reader.pages)
                if total_pages > MAX_PDF_PAGES:
                    raise ValueError(
                        f"Total page count exceeds the limit of {MAX_PDF_PAGES}."
                    )
                for page in reader.pages:
                    writer.add_page(page)

            output = io.BytesIO()
            writer.write(output)
            writer.close()
            output.seek(0)
            return _safe_send(output, "application/pdf", "merged.pdf")

        except ValueError as e:
            flash(str(e), "danger")
            return redirect(url_for("conversion.merge_pdf"))
        except Exception:
            logger.exception("PDF merge failed")
            flash("An error occurred while merging the PDF files.", "danger")
            return redirect(url_for("conversion.merge_pdf"))

    return render_template("doc_merge_pdf.html", form=form, max_files=MAX_MERGE_FILES)


# --- Split PDF ---
@conversion_bp.route("/convert/documents/split-pdf", methods=["GET", "POST"])
@limiter.limit("15 per minute;100 per hour", methods=["POST"])
def split_pdf():
    form = SplitPDFForm()

    if request.method == "POST":
        if not form.validate_on_submit():
            flash("Please select a PDF and provide a page range.", "danger")
            return redirect(url_for("conversion.split_pdf"))

        file = request.files.get("file")
        if not file or not file.filename:
            flash("No file selected.", "danger")
            return redirect(url_for("conversion.split_pdf"))

        try:
            from pypdf import PdfReader, PdfWriter

            data = _read_upload(file)
            if not _is_pdf_bytes(data):
                raise ValueError("Uploaded file is not a valid PDF.")
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ValueError("PDF is encrypted; please decrypt it first.")
            page_count = len(reader.pages)
            selected = _parse_page_ranges(form.page_ranges.data, page_count)

            writer = PdfWriter()
            for n in selected:
                writer.add_page(reader.pages[n - 1])

            output = io.BytesIO()
            writer.write(output)
            writer.close()
            output.seek(0)
            return _safe_send(output, "application/pdf", "split.pdf")

        except ValueError as e:
            flash(str(e), "danger")
            return redirect(url_for("conversion.split_pdf"))
        except Exception:
            logger.exception("PDF split failed")
            flash("An error occurred while splitting the PDF.", "danger")
            return redirect(url_for("conversion.split_pdf"))

    return render_template("doc_split_pdf.html", form=form)


# --- PDF -> TXT ---
@conversion_bp.route("/convert/documents/pdf-to-text", methods=["GET", "POST"])
@limiter.limit("15 per minute;100 per hour", methods=["POST"])
def pdf_to_text():
    form = SingleFileForm()

    if request.method == "POST":
        if not form.validate_on_submit():
            flash("Please select a PDF file.", "danger")
            return redirect(url_for("conversion.pdf_to_text"))

        file = request.files.get("file")
        if not file or not file.filename:
            flash("No file selected.", "danger")
            return redirect(url_for("conversion.pdf_to_text"))

        try:
            from pypdf import PdfReader

            data = _read_upload(file)
            if not _is_pdf_bytes(data):
                raise ValueError("Uploaded file is not a valid PDF.")
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise ValueError("PDF is encrypted; please decrypt it first.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ValueError(
                    f"Document exceeds the limit of {MAX_PDF_PAGES} pages."
                )

            chunks = []
            for page in reader.pages:
                chunks.append(page.extract_text() or "")
            text = "\n\n".join(chunks).strip() or "(no extractable text found)"

            output = io.BytesIO(text.encode("utf-8"))
            return _safe_send(output, "text/plain; charset=utf-8", "extracted.txt")

        except ValueError as e:
            flash(str(e), "danger")
            return redirect(url_for("conversion.pdf_to_text"))
        except Exception:
            logger.exception("PDF→TXT failed")
            flash("An error occurred while extracting text from the PDF.", "danger")
            return redirect(url_for("conversion.pdf_to_text"))

    return render_template("doc_pdf_to_text.html", form=form)


# --- DOCX -> TXT ---
@conversion_bp.route("/convert/documents/docx-to-text", methods=["GET", "POST"])
@limiter.limit("15 per minute;100 per hour", methods=["POST"])
def docx_to_text():
    form = SingleFileForm()

    if request.method == "POST":
        if not form.validate_on_submit():
            flash("Please select a DOCX file.", "danger")
            return redirect(url_for("conversion.docx_to_text"))

        file = request.files.get("file")
        if not file or not file.filename:
            flash("No file selected.", "danger")
            return redirect(url_for("conversion.docx_to_text"))

        try:
            data = _read_upload(file)
            if not _is_zip_bytes(data):
                raise ValueError("Uploaded file is not a valid DOCX document.")
            text = _docx_to_text(data) or "(empty document)"
            output = io.BytesIO(text.encode("utf-8"))
            return _safe_send(output, "text/plain; charset=utf-8", "extracted.txt")

        except ValueError as e:
            flash(str(e), "danger")
            return redirect(url_for("conversion.docx_to_text"))
        except Exception:
            logger.exception("DOCX→TXT failed")
            flash("An error occurred while reading the DOCX file.", "danger")
            return redirect(url_for("conversion.docx_to_text"))

    return render_template("doc_docx_to_text.html", form=form)


# --- DOCX -> PDF ---
@conversion_bp.route("/convert/documents/docx-to-pdf", methods=["GET", "POST"])
@limiter.limit("15 per minute;100 per hour", methods=["POST"])
def docx_to_pdf():
    form = SingleFileForm()

    if request.method == "POST":
        if not form.validate_on_submit():
            flash("Please select a DOCX file.", "danger")
            return redirect(url_for("conversion.docx_to_pdf"))

        file = request.files.get("file")
        if not file or not file.filename:
            flash("No file selected.", "danger")
            return redirect(url_for("conversion.docx_to_pdf"))

        try:
            data = _read_upload(file)
            if not _is_zip_bytes(data):
                raise ValueError("Uploaded file is not a valid DOCX document.")
            text = _docx_to_text(data) or "(empty document)"
            output = _text_to_pdf_stream(text)
            return _safe_send(output, "application/pdf", "converted.pdf")

        except ValueError as e:
            flash(str(e), "danger")
            return redirect(url_for("conversion.docx_to_pdf"))
        except Exception:
            logger.exception("DOCX→PDF failed")
            flash("An error occurred while converting the DOCX to PDF.", "danger")
            return redirect(url_for("conversion.docx_to_pdf"))

    return render_template("doc_docx_to_pdf.html", form=form)


# --- TXT -> PDF ---
@conversion_bp.route("/convert/documents/text-to-pdf", methods=["GET", "POST"])
@limiter.limit("15 per minute;100 per hour", methods=["POST"])
def text_to_pdf():
    form = SingleFileForm()

    if request.method == "POST":
        if not form.validate_on_submit():
            flash("Please select a TXT file.", "danger")
            return redirect(url_for("conversion.text_to_pdf"))

        file = request.files.get("file")
        if not file or not file.filename:
            flash("No file selected.", "danger")
            return redirect(url_for("conversion.text_to_pdf"))

        try:
            data = _read_upload(file)
            text = None
            for enc in ("utf-8", "utf-8-sig", "cp1250", "latin-1"):
                try:
                    text = data.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                raise ValueError("Could not decode the text file (unsupported encoding).")
            output = _text_to_pdf_stream(text)
            return _safe_send(output, "application/pdf", "converted.pdf")

        except ValueError as e:
            flash(str(e), "danger")
            return redirect(url_for("conversion.text_to_pdf"))
        except Exception:
            logger.exception("TXT→PDF failed")
            flash("An error occurred while converting the text to PDF.", "danger")
            return redirect(url_for("conversion.text_to_pdf"))

    return render_template("doc_text_to_pdf.html", form=form)


# --- Archive Conversion (Placeholder) ---
@conversion_bp.route("/convert/archive", methods=["GET"])
def archive_conversion():
    """Displays the archive tools page (Placeholder)."""
    return render_template("archive_conversion.html")
